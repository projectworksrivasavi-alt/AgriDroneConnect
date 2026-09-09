from flask import Flask, render_template, request, session, send_from_directory, send_file, redirect
from bson.objectid import ObjectId
from bson.errors import InvalidId
from markupsafe import escape
import os
import json
from io import BytesIO
from werkzeug.utils import secure_filename
from werkzeug.security import generate_password_hash, check_password_hash
from config import Config
from mongodb import operators, bookings, farmers
import random
from markupsafe import escape
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.config.from_object(Config)
if not app.config.get("SECRET_KEY"):
    raise RuntimeError("SECRET_KEY must be set in the environment")

UPLOAD_FOLDER = "/tmp/uploads" if os.getenv("VERCEL") else "uploads"
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
os.makedirs(UPLOAD_FOLDER, exist_ok=True)


# ---------------- HELPERS ----------------

def alert_redirect(message, url):
    """Return a small HTML page that alerts a message then redirects.
    Escapes message/url so user-controlled content can't break out of the
    <script> block (XSS)."""
    safe_message = escape(message)
    safe_url = escape(url)
    return f"""
    <script>
    alert("{safe_message}");
    window.location="{safe_url}";
    </script>
    """


def save_upload(file_storage):
    """Securely save an uploaded file and return its stored filename.
    Returns "" if no file was provided."""
    if not file_storage or file_storage.filename == "":
        return ""
    filename = secure_filename(file_storage.filename)
    file_storage.save(os.path.join(app.config["UPLOAD_FOLDER"], filename))
    return filename


def mark_payment_paid(booking_id):
    bookings.update_one(
        {"_id": ObjectId(booking_id)},
        {"$set": {"payment_status": "Paid"}}
    )


def get_object_id_or_none(id_str):
    """Safely convert a string to ObjectId, returning None if invalid."""
    try:
        return ObjectId(id_str)
    except (InvalidId, TypeError):
        return None


def export_collection(sheet, documents):
    """Write MongoDB documents to a worksheet with readable flat values."""
    documents = list(documents)
    keys = sorted({key for document in documents for key in document})
    if not keys:
        sheet.append(["No records"])
        return

    sheet.append(keys)
    for document in documents:
        row = []
        for key in keys:
            value = document.get(key, "")
            if isinstance(value, (dict, list, tuple)):
                value = json.dumps(value, default=str)
            elif value is not None:
                value = str(value) if isinstance(value, ObjectId) else value
            row.append(value)
        sheet.append(row)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 40)
        sheet.column_dimensions[column_letter].width = width


@app.route('/uploads/<filename>')
def uploaded_file(filename):
    return send_from_directory(app.config["UPLOAD_FOLDER"], filename)


# ---------------- HOME ----------------

@app.route('/')
def home():
    return render_template('index1.html')


# ---------------- FARMER ----------------
@app.route("/language")
def language():
    return render_template("language_selection.html")

translations = {

    "English": {

        # Farmer Login
        "login": "Farmer Login",
        "mobile": "Mobile Number",
        "mobile_placeholder": "Enter Mobile Number",
        "send_otp": "Send OTP",
        "enter_otp": "Enter OTP",
        "otp_placeholder": "Enter OTP",
        "verify_otp": "Verify OTP",

        # Farmer Dashboard
        "welcome_farmer": "Welcome Farmer",
        "book_service": "Book Drone Service",
        "my_bookings": "My Bookings",

        # Table Headings
        "crop": "Crop",
        "village": "Village",
        "area": "Area",
        "booking_date": "Booking Date",
        "status": "Status",
        "before_photo": "Before Photo",
        "after_photo": "After Photo",
        "payment": "Payment",
        "feedback": "Feedback",

        # Booking Status
        "pending": "Pending",
        "accepted": "Accepted",
        "completed": "Completed",

        # Common Buttons
        "view": "View",
        "pay_now": "Pay Now",
        "paid": "Paid",
        "give_feedback": "Give Feedback",
        "feedback_submitted": "Feedback Submitted",

        # Payment Page
        "payment_gateway": "Payment Gateway",
        "farmer": "Farmer",
        "acres": "Acres",
        "amount": "Amount",
        "book_service": "Book Drone Service",
        "farmer_name": "Farmer Name",
        "farmer_name_placeholder": "Enter Farmer Name",
        "village": "Village",
        "village_placeholder": "Enter Village",
        "crop": "Crop",
        "crop_placeholder": "Enter Crop Name",
        "area": "Area",
        "area_placeholder": "Enter Area in Acres",
        "booking_date": "Booking Date",
        "select_location": "Select Farm Location",
        "submit_booking": "Submit Booking"
    },
    

    "Telugu": {

        # Farmer Login
        "login": "రైతు లాగిన్",
        "mobile": "మొబైల్ నంబర్",
        "mobile_placeholder": "మొబైల్ నంబర్ నమోదు చేయండి",
        "send_otp": "OTP పంపండి",
        "enter_otp": "OTP నమోదు చేయండి",
        "otp_placeholder": "OTP నమోదు చేయండి",
        "verify_otp": "OTP నిర్ధారించండి",

        # Farmer Dashboard
        "welcome_farmer": "రైతుకు స్వాగతం",
        "book_service": "డ్రోన్ సేవను బుక్ చేయండి",
        "my_bookings": "నా బుకింగ్‌లు",

        # Table Headings
        "crop": "పంట",
        "village": "గ్రామం",
        "area": "విస్తీర్ణం",
        "booking_date": "బుకింగ్ తేదీ",
        "status": "స్థితి",
        "before_photo": "ముందు ఫోటో",
        "after_photo": "తర్వాత ఫోటో",
        "payment": "చెల్లింపు",
        "feedback": "అభిప్రాయం",

        # Booking Status
        "pending": "పెండింగ్",
        "accepted": "ఆమోదించబడింది",
        "completed": "పూర్తయింది",

        # Common Buttons
        "view": "చూడండి",
        "pay_now": "ఇప్పుడు చెల్లించండి",
        "paid": "చెల్లించబడింది",
        "give_feedback": "అభిప్రాయం ఇవ్వండి",
        "feedback_submitted": "అభిప్రాయం సమర్పించబడింది",

        # Payment Page
        "payment_gateway": "చెల్లింపు గేట్‌వే",
        "farmer": "రైతు",
        "acres": "ఎకరాలు",
        "amount": "మొత్తం",
        "book_service": "డ్రోన్ సేవను బుక్ చేయండి",
        "farmer_name": "రైతు పేరు",
        "farmer_name_placeholder": "రైతు పేరు నమోదు చేయండి",
        "village": "గ్రామం",
        "village_placeholder": "గ్రామం నమోదు చేయండి",
        "crop": "పంట",
        "crop_placeholder": "పంట పేరు నమోదు చేయండి",
        "area": "విస్తీర్ణం",
        "area_placeholder": "ఎకరాల్లో విస్తీర్ణం నమోదు చేయండి",
        "booking_date": "బుకింగ్ తేదీ",
        "select_location": "పొలం స్థానాన్ని ఎంచుకోండి",
        "submit_booking": "బుకింగ్ చేయండి"

    }

}

@app.route("/farmer")
def farmer():
    lang = session.get("language", "English")
    return render_template(
        "farmer.html",
        t=translations.get(lang, translations["English"])
    )


@app.route("/set_language", methods=["POST"])
def set_language():
    session["language"] = request.form["language"]
    return redirect("/farmer")


@app.route('/send_otp', methods=['POST'])
def send_otp():

    mobile = request.form['mobile']
    safe_mobile = escape(mobile)

    # Generate random 6-digit OTP
    otp = random.randint(100000, 999999)

    # Store OTP in session
    session["otp"] = str(otp)

    # Get selected language
    lang = session.get("language", "English")

    return render_template(
        "otp.html",
        otp=otp,
        mobile=mobile,
        t = translations.get(lang, translations["English"])
    )
    return f"""

    <h2>{t['send_otp']}</h2>

    OTP sent to <b>{safe_mobile}</b><br><br>

    <form action="/verify_otp" method="POST">

        <input type="hidden" name="mobile" value="{safe_mobile}">

        {t['enter_otp']} :
        <input type="text" name="otp" required>

        <button type="submit">{t['verify_otp']}</button>

    </form>
    """

@app.route('/verify_otp', methods=['POST'])
def verify_otp():

    mobile = request.form['mobile']
    otp = request.form['otp']

    # Verify generated OTP
    if otp != session.get("otp"):
        return "<script>alert('Invalid OTP');window.location='/farmer';</script>"

    session["farmer_mobile"] = mobile

    booking_list = list(bookings.find({"mobile": mobile}))

    # Selected language
    lang = session.get("language", "English")

    return render_template(
        "farmer_dashboard.html",
        bookings=booking_list,
        t=translations.get(lang, translations["English"])
    )

#------------- FARMER DASHBOARD ----------------

@app.route('/farmer_dashboard')
def farmer_dashboard():

    if "farmer_mobile" not in session:
        return '''
        <script>
        alert("Please Login First");
        window.location="/farmer";
        </script>
        '''

    mobile = session["farmer_mobile"]
    lang = session.get("language", "English")

    booking_list = list(bookings.find({"mobile": mobile}))

    return render_template(
        "farmer_dashboard.html",
        bookings=booking_list,
        t=translations.get(lang, translations["English"])
    )

# ---------------- BOOKING ----------------
@app.route('/booking')
def booking():

    if "farmer_mobile" not in session:
        return "<script>alert('Please Login First');window.location='/farmer';</script>"

    lang = session.get("language", "English")

    return render_template(
        "booking.html",
        t=translations.get(lang, translations["English"])
    )

@app.route('/book_service', methods=['POST'])
def book_service():

    try:
        latitude = float(request.form.get('latitude'))
        longitude = float(request.form.get('longitude'))
    except (TypeError, ValueError):
        return alert_redirect("Please select your farm location on the map", "/booking")

    booking = {
        "farmer_name": request.form.get('farmer_name'),
        "mobile": request.form.get('mobile'),
        "village": request.form.get('village'),
        "crop": request.form.get('crop'),
        "area": request.form.get('area'),
        "booking_date": request.form.get('date'),

        "latitude": latitude,
        "longitude": longitude,

        "status": "Pending",

        # New fields
        "assigned_operator": "",
        "rejected_by": [],

        # Existing project fields
        "before_photo": "",
        "after_photo": "",
        "payment_status": "Pending",
        "feedback": "",
        "rating": ""
    }

    bookings.insert_one(booking)

    return alert_redirect("Booking Submitted Successfully", "/farmer")


# ---------------- COMPLETED BOOKING ----------------

@app.route('/complete_booking/<id>', methods=['POST'])
def complete_booking(id):

    booking_oid = get_object_id_or_none(id)
    if booking_oid is None:
        return "<h2>Invalid booking id.</h2>", 400

    before_name = save_upload(request.files.get('before_photo'))
    after_name = save_upload(request.files.get('after_photo'))

    bookings.update_one(
        {"_id": booking_oid},
        {
            "$set": {
                "status": "Completed",
                "before_photo": before_name,
                "after_photo": after_name,
                "payment_status": "Pending"
            }
        }
    )

    return alert_redirect("Booking Completed Successfully", "/operator")


@app.route("/payment/<id>")
def payment(id):
    booking_oid = get_object_id_or_none(id)
    booking = bookings.find_one({"_id": booking_oid})

    lang = session.get("language", "English")

    return render_template(
        "payment.html",
        booking=booking,
        t=translations.get(lang, translations["English"])
    )

    if booking_oid is None:
        return "<h2>Invalid booking id.</h2>", 400

    booking = bookings.find_one({"_id": booking_oid})
    lang = session.get("language", "English")

    return render_template("payment.html",
    booking=booking,
    t=translations.get(lang, translations["English"])
)

@app.route("/payment_success/<id>", methods=["POST"])
def payment_success(id):
    booking_oid = get_object_id_or_none(id)
    if booking_oid is None:
        return "<h2>Invalid booking id.</h2>", 400

    mark_payment_paid(id)
    return alert_redirect("Payment Successful", "/farmer_dashboard")


@app.route("/pay_now/<id>", methods=["POST"])
def pay_now(id):
    booking_oid = get_object_id_or_none(id)
    if booking_oid is None:
        return "<h2>Invalid booking id.</h2>", 400

    mark_payment_paid(id)
    return alert_redirect("Payment Successful", "/farmer_dashboard")


@app.route("/feedback/<id>")
def feedback(id):
    booking_oid = get_object_id_or_none(id)
    if booking_oid is None:
        return "<h2>Invalid booking id.</h2>", 400

    booking = bookings.find_one({"_id": booking_oid})
    lang = session.get("language", "English")

    return render_template(
        "feedback.html",
        booking=booking,
        t=translations.get(lang, translations["English"])
    )


@app.route("/submit_feedback/<id>", methods=["POST"])
def submit_feedback(id):

    booking_oid = get_object_id_or_none(id)
    if booking_oid is None:
        return "<h2>Invalid booking id.</h2>", 400

    rating = request.form["rating"]
    comments = request.form["comments"]

    bookings.update_one(
        {"_id": booking_oid},
        {
            "$set": {
                "rating": rating,
                "comments": comments
            }
        }
    )

    return alert_redirect("Feedback Submitted Successfully", "/farmer_dashboard")


# ---------------- OPERATOR LOGIN ----------------

@app.route('/operator')
def operator():
    return render_template('operator_login.html')


@app.route('/operator_register_page')
def operator_register_page():
    return render_template('operator_register.html')


@app.route('/operator_register', methods=['POST'])
def operator_register():

    # Personal Details
    name = request.form['name']
    mobile = request.form['mobile']
    email = request.form['email']
    password = request.form['password']
    confirm_password = request.form['confirm_password']

    if password != confirm_password:
        return "<h2>Passwords do not match.</h2>"

    aadhaar = request.form['aadhaar']
    license_no = request.form['license']
    drone_no = request.form['drone_no']

    # Uploaded files - saved via shared helper (handles empty/optional files too)
    photo_filename = save_upload(request.files.get('photo'))
    aadhaar_filename = save_upload(request.files.get('aadhaar_file'))
    license_filename = save_upload(request.files.get('license_file'))
    registration_filename = save_upload(request.files.get('registration_file'))
    insurance_filename = save_upload(request.files.get('insurance_file'))

    # Store only a one-way password hash.
    operator = {
        "name": name,
        "mobile": mobile,
        "email": email,
        "password_hash": generate_password_hash(password),
        "aadhaar": aadhaar,
        "license_no": license_no,
        "drone_no": drone_no,
        "photo": photo_filename,
        "aadhaar_file": aadhaar_filename,
        "license_file": license_filename,
        "registration_file": registration_filename,
        "insurance_file": insurance_filename,
        "status": "Pending"
    }

    operators.insert_one(operator)

    return "<h2>Registration Successful.<br>Waiting for Admin Approval.</h2>"


@app.route('/operator_login', methods=['POST'])
def operator_login():

    mobile = request.form['mobile']
    password = request.form['password']

    operator = operators.find_one({"mobile": mobile})

    if not operator or not check_password_hash(
        operator.get("password_hash", ""), password
    ):
        return "<h2>Invalid Mobile Number or Password.</h2>"

    if operator["status"] != "Approved":
        return "<h2>Your account is waiting for Admin Approval.</h2>"

    session["operator_mobile"] = mobile

    booking_list = list(bookings.find({
        "$or": [
            {
                "status": "Pending",
                "rejected_by": {
                    "$ne": session["operator_mobile"]
                }
            },
            {
                "assigned_operator": session["operator_mobile"]
            }
        ]
    }))
    return render_template(
        "operator_dashboard.html",
        operator=operator,
        bookings=booking_list
    )


@app.route('/approve_operator/<id>', methods=['POST'])
def approve_operator(id):
    operator_oid = get_object_id_or_none(id)
    if operator_oid is None:
        return "<h2>Invalid operator id.</h2>", 400

    operators.update_one(
        {"_id": operator_oid},
        {"$set": {"status": "Approved"}}
    )
    return alert_redirect("Operator Approved Successfully", "/admin")


@app.route('/reject_operator/<id>', methods=['POST'])
def reject_operator(id):
    operator_oid = get_object_id_or_none(id)
    if operator_oid is None:
        return "<h2>Invalid operator id.</h2>", 400

    operators.delete_one({"_id": operator_oid})
    return alert_redirect("Operator Rejected", "/admin")


# ---------------- ACCEPT / REJECT BOOKING ----------------

@app.route('/accept_booking/<id>', methods=['POST'])
def accept_booking(id):

    if "operator_mobile" not in session:
        return '''
        <script>
        alert("Please Login First");
        window.location="/operator";
        </script>
        '''

    booking_oid = get_object_id_or_none(id)
    if booking_oid is None:
        return "<h2>Invalid booking id.</h2>", 400

    bookings.update_one(
        {"_id": booking_oid},
        {
            "$set": {
                "status": "Accepted",
                "assigned_operator": session["operator_mobile"]
            }
        }
    )
    return '''
    <script>
    alert("Booking Accepted Successfully");
    window.history.back();
    </script>
    '''


@app.route('/reject_booking/<id>', methods=['POST'])
def reject_booking(id):

    if "operator_mobile" not in session:
        return '''
        <script>
        alert("Please Login First");
        window.location="/operator";
        </script>
        '''

    booking_oid = get_object_id_or_none(id)
    if booking_oid is None:
        return "<h2>Invalid booking id.</h2>", 400

    bookings.update_one(
        {"_id": booking_oid},
        {
            "$push": {
                "rejected_by": session["operator_mobile"]
            }
        }
    )

    return '''
    <script>
    alert("Booking Rejected");
    window.history.back();
    </script>
    '''


# ---------------- ADMIN ----------------

@app.route('/admin')
def admin():
    return render_template('admin.html')


@app.route('/admin_login', methods=['POST'])
def admin_login():

    username = request.form['username']
    password = request.form['password']

    admin_username = os.getenv("ADMIN_USERNAME")
    admin_password_hash = os.getenv("ADMIN_PASSWORD_HASH")

    if (
        admin_username
        and admin_password_hash
        and username == admin_username
        and check_password_hash(admin_password_hash, password)
    ):

        session["admin_logged_in"] = True

        operator_list = list(operators.find())
        farmer_list = list(bookings.find())   # Farmer details from bookings

        return render_template(
            "admin_dashboard.html",
            operators=operator_list,
            farmers=farmer_list,
            total_farmers=len(farmer_list),
            total_operators=len(operator_list),
            total_bookings=bookings.count_documents({}),
            pending=bookings.count_documents({"status": "Pending"}),
            accepted=bookings.count_documents({"status": "Accepted"}),
            completed=bookings.count_documents({"status": "Completed"}),
            rejected=bookings.count_documents({"status": "Rejected"})
        )

    return "<h2>Invalid Username or Password</h2>"


@app.route('/admin/export_excel')
def export_excel():
    if not session.get("admin_logged_in"):
        return alert_redirect("Please login as admin first", "/admin")

    workbook = Workbook()
    workbook.remove(workbook.active)
    export_collection(workbook.create_sheet("Bookings"), bookings.find())
    export_collection(workbook.create_sheet("Farmers"), farmers.find())
    export_collection(workbook.create_sheet("Operators"), operators.find())

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="agridrone_mongodb_data.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.route('/logout')
def logout():
    session.clear()
    return alert_redirect("Logged Out Successfully", "/")


# ---------------- RUN ----------------

if __name__ == '__main__':
    app.run(host="127.0.0.1", port=5000, debug=False)